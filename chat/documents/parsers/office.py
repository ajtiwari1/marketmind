"""Word and PowerPoint parsers preserving their structural boundaries."""

from pathlib import Path

from docx import Document as WordDocument
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..models import DocumentSection


def parse_docx(path: Path) -> tuple[DocumentSection, ...]:
    document = WordDocument(path)
    sections: list[DocumentSection] = []
    heading_path: list[str] = []
    body = document.element.body
    for element in body.iterchildren():
        if element.tag.endswith("}p"):
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            if text:
                style_name = paragraph.style.name if paragraph.style else ""
                if style_name.startswith("Heading"):
                    try:
                        depth = int(style_name.removeprefix("Heading").strip())
                    except ValueError:
                        depth = 1
                    heading_path[:] = heading_path[: max(0, depth - 1)]
                    heading_path.append(text)
                    continue
                sections.append(
                    DocumentSection(
                        text=text,
                        metadata={
                            "section_index": len(sections),
                            "heading_path": tuple(heading_path),
                            "style": style_name,
                        },
                    )
                )
        elif element.tag.endswith("}tbl"):
            table = WordTable(element, document)
            rows = [
                " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells)
                for row in table.rows
            ]
            if rows:
                sections.append(
                    DocumentSection(
                        text="\n".join(rows),
                        kind="table",
                        metadata={
                            "section_index": len(sections),
                            "heading_path": tuple(heading_path),
                        },
                    )
                )
    return tuple(sections)


def _shape_text(shape: object) -> list[DocumentSection]:
    sections: list[DocumentSection] = []
    if getattr(shape, "has_text_frame", False):
        text = shape.text.strip()
        if text:
            sections.append(DocumentSection(text=text))
    if getattr(shape, "has_table", False):
        table = shape.table
        rows = [
            " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells)
            for row in table.rows
        ]
        if rows:
            sections.append(DocumentSection(text="\n".join(rows), kind="table"))
    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            sections.extend(_shape_text(child))
    return sections


def parse_pptx(path: Path) -> tuple[DocumentSection, ...]:
    presentation = Presentation(path)
    sections: list[DocumentSection] = []
    for slide_index, slide in enumerate(presentation.slides):
        title = slide.shapes.title.text.strip() if slide.shapes.title else None
        for shape in slide.shapes:
            for parsed in _shape_text(shape):
                metadata = dict(parsed.metadata)
                metadata["slide_number"] = slide_index + 1
                if title:
                    metadata["slide_title"] = title
                sections.append(
                    DocumentSection(
                        text=parsed.text,
                        kind=parsed.kind,
                        metadata=metadata,
                    )
                )
    return tuple(sections)
