"""CSV, Excel, and JSON parsers."""

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..models import DocumentSection


def _row_text(row: tuple[Any, ...] | list[Any]) -> str:
    return " | ".join("" if value is None else str(value).strip() for value in row)


def parse_csv(path: Path) -> tuple[DocumentSection, ...]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        rows = [row for row in csv.reader(file) if any(cell.strip() for cell in row)]
    if not rows:
        return ()
    header = _row_text(rows[0])
    content = "\n".join([f"Columns: {header}", *(_row_text(row) for row in rows[1:])])
    return (
        DocumentSection(
            text=content,
            kind="table",
            metadata={"table_name": path.stem, "row_start": 2},
        ),
    )


def parse_xlsx(path: Path) -> tuple[DocumentSection, ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[DocumentSection] = []
    try:
        for worksheet in workbook.worksheets:
            nonempty_rows = [
                tuple(row)
                for row in worksheet.iter_rows(values_only=True)
                if any(value is not None and str(value).strip() for value in row)
            ]
            if not nonempty_rows:
                continue
            header = _row_text(nonempty_rows[0])
            content = "\n".join(
                [
                    f"Sheet: {worksheet.title} | Columns: {header}",
                    *(_row_text(row) for row in nonempty_rows[1:]),
                ]
            )
            sections.append(
                DocumentSection(
                    text=content,
                    kind="table",
                    metadata={"sheet_name": worksheet.title, "row_start": 2},
                )
            )
    finally:
        workbook.close()
    return tuple(sections)


def _json_sections(value: Any, key_path: str = "$") -> list[DocumentSection]:
    if isinstance(value, dict):
        sections: list[DocumentSection] = []
        for key, child in value.items():
            sections.extend(_json_sections(child, f"{key_path}.{key}"))
        if not sections:
            sections.append(
                DocumentSection(
                    text=f"{key_path}: {{}}",
                    metadata={"key_path": key_path},
                )
            )
        return sections
    if isinstance(value, list):
        sections = []
        for index, child in enumerate(value):
            item_path = f"{key_path}[{index}]"
            if isinstance(child, (dict, list)):
                sections.append(
                    DocumentSection(
                        text=f"{item_path}: {json.dumps(child, ensure_ascii=False)}",
                        metadata={"key_path": item_path, "record_index": index},
                    )
                )
            else:
                sections.extend(_json_sections(child, item_path))
        if not sections:
            sections.append(
                DocumentSection(
                    text=f"{key_path}: []",
                    metadata={"key_path": key_path},
                )
            )
        return sections
    return [
        DocumentSection(
            text=f"{key_path}: {json.dumps(value, ensure_ascii=False)}",
            metadata={"key_path": key_path},
        )
    ]


def parse_json(path: Path) -> tuple[DocumentSection, ...]:
    with path.open("r", encoding="utf-8-sig") as file:
        value = json.load(file)
    return tuple(_json_sections(value))
