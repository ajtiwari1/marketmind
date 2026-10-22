import csv
import json
import math
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

import pymupdf as fitz
from docx import Document as WordDocument
from openpyxl import Workbook
from pptx import Presentation

from chat.chunking.text import chunk_document
from chat.config import RagConfig
from chat.documents.loader import DocumentLoadError, load_document
from chat.documents.models import Document, DocumentSection
from chat.embeddings.sparse import SparseVector
from chat.pipeline import IngestionError, ingest_file


class CharacterTokenizer:
    def encode(self, text: str, **kwargs: object) -> list[int]:
        del kwargs
        return list(text.encode("utf-8"))

    def decode(self, token_ids: list[int], **kwargs: object) -> str:
        del kwargs
        return bytes(token_ids).decode("utf-8", errors="ignore")


class FakeEmbedder:
    def __init__(self) -> None:
        self.tokenizer = CharacterTokenizer()
        self.dimension = 3
        self.model_name = "fake-model"
        self.max_input_tokens = 1000
        self.encoded_texts: list[str] = []

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        self.encoded_texts.extend(texts)
        return [[0.1, 0.2, 0.3] for _ in texts]


class FakeSparseEmbedder:
    def __init__(self) -> None:
        self.model_name = "fake-sparse-model"
        self.encoded_texts: list[str] = []

    def embed_documents(self, texts: list[str]) -> list[SparseVector]:
        self.encoded_texts.extend(texts)
        return [SparseVector(indices=[1, 4], values=[0.6, 0.4]) for _ in texts]


class NonFiniteEmbedder(FakeEmbedder):
    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [[math.nan, 0.2, 0.3] for _ in texts]


class DocumentLoaderTests(unittest.TestCase):
    def test_loads_text_and_generates_stable_document_id(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "guide.txt")
            path.write_text("First paragraph.\n\nSecond paragraph.", encoding="utf-8")

            first = load_document(path)
            second = load_document(path)

        self.assertEqual(first.document_id, second.document_id)
        self.assertEqual(len(first.content_hash), 64)
        self.assertEqual([section.text for section in first.sections], [
            "First paragraph.",
            "Second paragraph.",
        ])

    def test_markdown_preserves_heading_path_and_code(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "guide.md")
            path.write_text(
                "# Setup\n\nInstall the package.\n\n```py\nprint('ok')\n```",
                encoding="utf-8",
            )

            document = load_document(path)

        self.assertEqual(document.sections[0].metadata["heading_path"], ("Setup",))
        self.assertIn("Install the package.", document.sections[0].text)
        self.assertEqual(document.sections[1].kind, "code")

    def test_loads_supported_formats_with_provenance(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)

            pdf_path = root / "guide.pdf"
            pdf = fitz.open()
            page = pdf.new_page()
            page.insert_text((72, 72), "PDF content")
            pdf.save(pdf_path)
            pdf.close()

            docx_path = root / "guide.docx"
            word = WordDocument()
            word.add_heading("Heading", level=1)
            word.add_paragraph("Word content")
            word.add_table(rows=1, cols=2).cell(0, 0).text = "table cell"
            word.save(docx_path)

            pptx_path = root / "slides.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = "Slide title"
            slide.shapes.add_textbox(100, 100, 300, 100).text = "Slide content"
            presentation.save(pptx_path)

            xlsx_path = root / "sheet.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Inventory"
            sheet.append(["Name", "Count"])
            sheet.append(["Pencils", 12])
            workbook.save(xlsx_path)

            csv_path = root / "records.csv"
            with csv_path.open("w", newline="", encoding="utf-8") as file:
                csv.writer(file).writerows([["Name", "Count"], ["Books", "8"]])

            html_path = root / "page.html"
            html_path.write_text(
                "<html><head><title>Page</title></head><body><nav>Noise</nav>"
                "<main><h1>Section</h1><p>HTML content</p></main></body></html>",
                encoding="utf-8",
            )

            json_path = root / "data.json"
            json_path.write_text(json.dumps({"items": [{"name": "value"}]}))

            documents = {
                suffix: load_document(root / filename)
                for suffix, filename in {
                    "pdf": "guide.pdf",
                    "docx": "guide.docx",
                    "pptx": "slides.pptx",
                    "xlsx": "sheet.xlsx",
                    "csv": "records.csv",
                    "html": "page.html",
                    "json": "data.json",
                }.items()
            }

        self.assertEqual(documents["pdf"].sections[0].metadata["page_number"], 1)
        self.assertTrue(any("Word content" in section.text for section in documents["docx"].sections))
        self.assertTrue(any(
            section.metadata["heading_path"] == ("Heading",)
            for section in documents["docx"].sections
        ))
        self.assertTrue(any(section.kind == "table" for section in documents["docx"].sections))
        self.assertTrue(any(section.metadata["slide_number"] == 1 for section in documents["pptx"].sections))
        self.assertEqual(documents["xlsx"].sections[0].metadata["sheet_name"], "Inventory")
        self.assertEqual(documents["csv"].sections[0].metadata["row_start"], 2)
        self.assertNotIn("Noise", " ".join(section.text for section in documents["html"].sections))
        self.assertIn("$.items[0]", documents["json"].sections[0].text)

    def test_rejects_unsupported_and_scanned_documents(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            unsupported = Path(directory, "archive.bin")
            unsupported.write_bytes(b"\x00\x01")
            empty_pdf = Path(directory, "scan.pdf")
            pdf = fitz.open()
            pdf.new_page()
            pdf.save(empty_pdf)
            pdf.close()

            with self.assertRaisesRegex(DocumentLoadError, "Unsupported document type"):
                load_document(unsupported)
            with self.assertRaisesRegex(DocumentLoadError, "scanned PDFs require OCR"):
                load_document(empty_pdf)

    def test_enforces_file_size_limit(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "large.txt")
            path.write_text("content", encoding="utf-8")

            with self.assertRaisesRegex(DocumentLoadError, "size limit"):
                load_document(path, max_file_size_bytes=2)


class ChunkingTests(unittest.TestCase):
    def test_chunks_oversized_sections_and_preserves_heading_metadata(self) -> None:
        text = " ".join(["word"] * 100)
        document = Document(
            document_id="doc-id",
            source="file:///guide.md",
            file_name="guide.md",
            file_type="text/markdown",
            content_hash="hash",
            sections=(
                DocumentSection(
                    text=text,
                    metadata={"heading_path": ("Guide", "Install")},
                ),
            ),
        )

        chunks = chunk_document(
            document,
            CharacterTokenizer(),
            max_tokens=100,
            overlap_tokens=10,
        )

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk.text.encode("utf-8")) <= 100 for chunk in chunks))
        self.assertTrue(all("Guide > Install" in chunk.text for chunk in chunks))
        self.assertEqual([chunk.chunk_index for chunk in chunks], list(range(len(chunks))))
        self.assertEqual(len({chunk.chunk_id for chunk in chunks}), len(chunks))
        previous_body = chunks[0].text.split("\n", 1)[1]
        next_body = chunks[1].text.split("\n", 1)[1]
        self.assertTrue(next_body.startswith(previous_body[-10:].strip()))

    def test_table_chunks_repeat_headers(self) -> None:
        document = Document(
            document_id="doc-id",
            source="file:///items.csv",
            file_name="items.csv",
            file_type="text/csv",
            content_hash="hash",
            sections=(
                DocumentSection(
                    text="Name | Count\n" + "\n".join(f"Item {i} | {i}" for i in range(20)),
                    kind="table",
                ),
            ),
        )

        chunks = chunk_document(document, CharacterTokenizer(), max_tokens=60, overlap_tokens=0)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.text.startswith("Name | Count") for chunk in chunks))
        self.assertTrue(all(len(chunk.text.encode("utf-8")) <= 60 for chunk in chunks))


class PipelineTests(unittest.TestCase):
    def setUp(self) -> None:
        client = MagicMock()
        patcher = patch("chat.pipeline.get_langfuse_client", return_value=client)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.langfuse_client = client

    def test_dry_run_does_not_embed_or_write_to_qdrant(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "guide.txt")
            path.write_text("A short document.", encoding="utf-8")
            embedder = FakeEmbedder()
            store = MagicMock()
            config = RagConfig(
                embedding_model="fake-model",
                sparse_embedding_model="fake-sparse-model",
                chunk_size_tokens=100,
                chunk_overlap_tokens=0,
            )

            summary = ingest_file(
                path,
                config=config,
                dry_run=True,
                embedder=embedder,
                sparse_embedder=FakeSparseEmbedder(),
                store=store,
            )

        self.assertEqual(summary.accepted_files, 1)
        self.assertEqual(summary.sections_extracted, 1)
        self.assertEqual(summary.chunks_upserted, 0)
        self.assertEqual(embedder.encoded_texts, [])
        store.ensure_collection.assert_not_called()
        store.upsert_document.assert_not_called()

    def test_ingest_embeds_and_upserts_chunks(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "guide.txt")
            path.write_text("A short document.", encoding="utf-8")
            embedder = FakeEmbedder()
            embedder.max_input_tokens = 16
            sparse_embedder = FakeSparseEmbedder()
            store = MagicMock()
            store.upsert_document.return_value = 1
            config = RagConfig(
                embedding_model="fake-model",
                sparse_embedding_model="fake-sparse-model",
                chunk_size_tokens=100,
                chunk_overlap_tokens=0,
            )

            summary = ingest_file(
                path,
                config=config,
                create_collection=True,
                embedder=embedder,
                sparse_embedder=sparse_embedder,
                store=store,
            )

        self.assertEqual(summary.chunks_upserted, 1)
        self.assertGreater(len(embedder.encoded_texts), 1)
        self.assertEqual(embedder.encoded_texts, sparse_embedder.encoded_texts)
        store.ensure_collection.assert_called_once_with(3, create_if_missing=True)
        store.upsert_document.assert_called_once()
        stored_chunks = store.upsert_document.call_args.args[0]
        self.assertTrue(all(len(chunk.text.encode("utf-8")) <= 16 for chunk in stored_chunks))
        self.assertEqual(len(store.upsert_document.call_args.args[1]), len(stored_chunks))
        self.assertEqual(len(store.upsert_document.call_args.args[2]), len(stored_chunks))
        self.assertEqual(
            store.upsert_document.call_args.args[2][0],
            SparseVector(indices=[1, 4], values=[0.6, 0.4]),
        )
        trace_names = [
            call.kwargs["name"]
            for call in self.langfuse_client.start_as_current_observation.call_args_list
        ]
        self.assertEqual(
            trace_names,
            [
                "rag.ingest_document",
                "rag.embedding.dense",
                "rag.embedding.sparse",
                "rag.qdrant.upsert",
            ],
        )
        trace_inputs = [
            call.kwargs.get("input")
            for call in self.langfuse_client.start_as_current_observation.call_args_list
        ]
        self.assertTrue(all("A short document." not in str(value) for value in trace_inputs))

    def test_rejects_non_finite_embeddings_before_qdrant_write(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "guide.txt")
            path.write_text("A short document.", encoding="utf-8")
            embedder = NonFiniteEmbedder()
            store = MagicMock()
            config = RagConfig(
                embedding_model="fake-model",
                sparse_embedding_model="fake-sparse-model",
                chunk_size_tokens=100,
                chunk_overlap_tokens=0,
            )

            with self.assertRaisesRegex(IngestionError, "non-finite values"):
                ingest_file(
                    path,
                    config=config,
                    embedder=embedder,
                    sparse_embedder=FakeSparseEmbedder(),
                    store=store,
                )

        store.upsert_document.assert_not_called()


if __name__ == "__main__":
    unittest.main()
